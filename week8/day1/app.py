import uuid
import numbers
import io

import pandas as pd
import pandasai as pai
import plotly.express as px
import ollama

from fastapi import FastAPI, UploadFile, File, Form
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

from pandasai_litellm.litellm import LiteLLM

from pptx import Presentation


# ==================================================
# APP
# ==================================================

app = FastAPI(title="CIA CSV Analyst")

# Temporary CSV storage
dataframes = {}


# ==================================================
# FRONTEND
# ==================================================

app.mount(
    "/static",
    StaticFiles(directory="static"),
    name="static"
)


@app.get("/")
def home():
    return FileResponse("static/index.html")


# ==================================================
# FORMAT VALUE
# ==================================================

def format_value(value):

    if value is None:
        return "No result found."

    if isinstance(value, numbers.Number):

        if pd.isna(value):
            return "No result found."

        if float(value).is_integer():
            return f"{int(value):,}"

        return f"{float(value):,.2f}"

    return str(value)


# ==================================================
# FIND COLUMN
# ==================================================

def find_column(df, possible_names):

    normalized = {
        str(column)
        .lower()
        .replace("_", " ")
        .strip(): column
        for column in df.columns
    }

    for name in possible_names:

        name = (
            name
            .lower()
            .replace("_", " ")
            .strip()
        )

        if name in normalized:
            return normalized[name]

    return None


# ==================================================
# CSV UPLOAD
# ==================================================

@app.post("/upload")
async def upload_csv(file: UploadFile = File(...)):

    try:

        if not file.filename.lower().endswith(".csv"):

            return {
                "success": False,
                "error": "Please upload a CSV file."
            }

        df = pd.read_csv(file.file)

        session_id = str(uuid.uuid4())

        dataframes[session_id] = df

        return {
            "success": True,
            "session_id": session_id,
            "filename": file.filename,
            "rows": len(df),
            "columns": list(df.columns)
        }

    except Exception as e:

        return {
            "success": False,
            "error": str(e)
        }


# ==================================================
# CIA QUERY
# ==================================================

@app.post("/cia/sql-analyst")
async def cia_query(
    session_id: str = Form(...),
    question: str = Form(...)
):

    try:

        if session_id not in dataframes:

            return {
                "success": False,
                "error": "Session not found. Please upload the CSV again."
            }

        df = dataframes[session_id]

        # ------------------------------------------
        # LOCAL OLLAMA
        # ------------------------------------------

        llm = LiteLLM(
            model="ollama/gemma3:latest",
            api_base="http://localhost:11434"
        )

        pai.config.set({
            "llm": llm
        })

        # PandasAI DataFrame
        pandasai_df = pai.DataFrame(df)

        # Ask question
        result = pandasai_df.chat(question)

        # ------------------------------------------
        # CLEAN ANSWER
        # ------------------------------------------

        if isinstance(result, pd.DataFrame):

            if result.shape[0] == 1 and result.shape[1] == 1:

                answer = format_value(
                    result.iloc[0, 0]
                )

            else:

                answer = result.to_string(
                    index=False
                )

        elif isinstance(result, pd.Series):

            if len(result) == 1:

                answer = format_value(
                    result.iloc[0]
                )

            else:

                answer = result.to_string()

        else:

            answer = format_value(result)

        # ------------------------------------------
        # CHART
        # ------------------------------------------

        chart = create_chart(
            df,
            question
        )

        return {
            "success": True,
            "question": question,
            "answer": answer,
            "chart": chart
        }

    except Exception as e:

        return {
            "success": False,
            "error": str(e)
        }


# ==================================================
# CREATE CHART
# ==================================================

def create_chart(df, question):

    q = question.lower()

    category_col = find_column(
        df,
        [
            "category",
            "product_category"
        ]
    )

    city_col = find_column(
        df,
        [
            "city"
        ]
    )

    sales_col = find_column(
        df,
        [
            "sales",
            "sale",
            "amount",
            "revenue",
            "booking_value",
            "sales_inr"
        ]
    )

    payment_col = find_column(
        df,
        [
            "payment_method",
            "payment method",
            "payment"
        ]
    )

    # ------------------------------------------
    # CATEGORY SALES
    # ------------------------------------------

    if (
        "category" in q
        and "sales" in q
    ):

        if category_col and sales_col:

            chart_df = (
                df.groupby(category_col)[sales_col]
                .sum()
                .reset_index()
                .sort_values(
                    sales_col,
                    ascending=False
                )
            )

            fig = px.bar(
                chart_df,
                x=category_col,
                y=sales_col,
                title="Sales by Category"
            )

            return fig.to_json()

    # ------------------------------------------
    # CITY SALES
    # ------------------------------------------

    if (
        "city" in q
        and "sales" in q
    ):

        if city_col and sales_col:

            chart_df = (
                df.groupby(city_col)[sales_col]
                .sum()
                .reset_index()
                .sort_values(
                    sales_col,
                    ascending=False
                )
                .head(10)
            )

            fig = px.bar(
                chart_df,
                x=city_col,
                y=sales_col,
                title="Top Cities by Sales"
            )

            return fig.to_json()

    # ------------------------------------------
    # PAYMENT METHOD
    # ------------------------------------------

    if "payment" in q:

        if payment_col:

            chart_df = (
                df[payment_col]
                .value_counts()
                .reset_index()
            )

            chart_df.columns = [
                payment_col,
                "count"
            ]

            fig = px.bar(
                chart_df,
                x=payment_col,
                y="count",
                title="Payment Method Usage"
            )

            return fig.to_json()

    return None


# ==================================================
# GENERATE 5 EDA INSIGHTS
# ==================================================

def generate_insights(df):

    insights = []

    # ------------------------------------------
    # FIND COLUMNS
    # ------------------------------------------

    sales_col = find_column(
        df,
        [
            "sales",
            "sales_inr",
            "sale",
            "revenue",
            "amount",
            "booking_value"
        ]
    )

    profit_col = find_column(
        df,
        [
            "profit",
            "profit_inr"
        ]
    )

    category_col = find_column(
        df,
        [
            "category",
            "product_category"
        ]
    )

    city_col = find_column(
        df,
        [
            "city"
        ]
    )

    payment_col = find_column(
        df,
        [
            "payment_method",
            "payment method",
            "payment"
        ]
    )

    # ------------------------------------------
    # INSIGHT 1
    # DATASET SIZE
    # ------------------------------------------

    insights.append(
        f"The dataset contains "
        f"{len(df):,} records and "
        f"{len(df.columns)} columns."
    )

    # ------------------------------------------
    # INSIGHT 2
    # TOTAL SALES
    # ------------------------------------------

    if sales_col:

        sales_values = pd.to_numeric(
            df[sales_col],
            errors="coerce"
        )

        total_sales = sales_values.sum()

        insights.append(
            f"Total sales are "
            f"{total_sales:,.2f}."
        )

    else:

        insights.append(
            "A sales column was not available "
            "for total sales analysis."
        )

    # ------------------------------------------
    # INSIGHT 3
    # TOP CATEGORY
    # ------------------------------------------

    if category_col and sales_col:

        category_sales = (
            df.groupby(category_col)[sales_col]
            .sum()
            .sort_values(
                ascending=False
            )
        )

        if not category_sales.empty:

            top_category = (
                category_sales.index[0]
            )

            top_category_sales = (
                category_sales.iloc[0]
            )

            insights.append(
                f"{top_category} is the "
                f"highest-selling category "
                f"with sales of "
                f"{top_category_sales:,.2f}."
            )

        else:

            insights.append(
                "No category sales data "
                "was available."
            )

    else:

        insights.append(
            "Category sales analysis "
            "was not available."
        )

    # ------------------------------------------
    # INSIGHT 4
    # TOP CITY
    # ------------------------------------------

    if city_col and sales_col:

        city_sales = (
            df.groupby(city_col)[sales_col]
            .sum()
            .sort_values(
                ascending=False
            )
        )

        if not city_sales.empty:

            top_city = city_sales.index[0]

            top_city_sales = city_sales.iloc[0]

            insights.append(
                f"{top_city} is the "
                f"highest-performing city "
                f"with sales of "
                f"{top_city_sales:,.2f}."
            )

        else:

            insights.append(
                "No city sales data "
                "was available."
            )

    else:

        insights.append(
            "City sales analysis "
            "was not available."
        )

    # ------------------------------------------
    # INSIGHT 5
    # PROFIT OR PAYMENT
    # ------------------------------------------

    if profit_col:

        profit_values = pd.to_numeric(
            df[profit_col],
            errors="coerce"
        )

        average_profit = (
            profit_values.mean()
        )

        insights.append(
            f"Average profit is "
            f"{average_profit:,.2f}."
        )

    elif payment_col:

        payment_counts = (
            df[payment_col]
            .value_counts()
        )

        if not payment_counts.empty:

            top_payment = (
                payment_counts.index[0]
            )

            top_payment_count = (
                payment_counts.iloc[0]
            )

            insights.append(
                f"{top_payment} is the most "
                f"frequently used payment method "
                f"with {top_payment_count:,} transactions."
            )

        else:

            insights.append(
                "No payment method data "
                "was available."
            )

    else:

        insights.append(
            "Profit and payment information "
            "was not available."
        )

    # ------------------------------------------
    # EXACTLY 5 INSIGHTS
    # ------------------------------------------

    return insights[:5]


# ==================================================
# LOCAL OLLAMA AI SUMMARY
# ==================================================

# ==================================================
# LOCAL OLLAMA AI SUMMARY
# ==================================================

def generate_ai_summary(insights):

    prompt = f"""
You are a business data analyst.

Create a professional executive summary using
the following five data insights.

Requirements:

- Maximum 200 words
- Use professional business language
- Clearly explain the important findings
- Mention important performance trends
- Do not invent any information
- End with a short actionable recommendation
- Return only the executive summary
- Do not use markdown headings

DATA INSIGHTS:

1. {insights[0]}

2. {insights[1]}

3. {insights[2]}

4. {insights[3]}

5. {insights[4]}
"""

    response = ollama.chat(
        model="gemma3:latest",
        messages=[
            {
                "role": "user",
                "content": prompt
            }
        ]
    )

    return response["message"]["content"]

# ==================================================
# AI SUMMARY ENDPOINT
# ==================================================

@app.post("/ai-summary")
async def ai_summary(
    session_id: str = Form(...)
):

    try:

        if session_id not in dataframes:

            return {
                "success": False,
                "error": "Session not found. Please upload the CSV again."
            }

        df = dataframes[session_id]

        # Generate 5 EDA insights
        insights = generate_insights(df)

        # Generate AI summary
        summary = generate_ai_summary(
            insights
        )

        return {
            "success": True,
            "insights": insights,
            "summary": summary
        }

    except Exception as e:

        return {
            "success": False,
            "error": str(e)
        }


# ==================================================
# CREATE POWERPOINT
# ==================================================

def create_ppt(
    insights,
    summary
):

    prs = Presentation()

    # ------------------------------------------
    # SLIDE 1 - TITLE
    # ------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[0]
    )

    slide.shapes.title.text = (
        "CIA Data Analysis Report"
    )

    slide.placeholders[1].text = (
        "AI-Powered Executive Summary"
    )

    # ------------------------------------------
    # SLIDES 2-6 - FIVE INSIGHTS
    # ------------------------------------------

    for index, insight in enumerate(
        insights,
        start=1
    ):

        slide = prs.slides.add_slide(
            prs.slide_layouts[1]
        )

        slide.shapes.title.text = (
            f"Key Insight {index}"
        )

        slide.placeholders[1].text = insight

    # ------------------------------------------
    # SLIDE 7 - RECOMMENDATION
    # ------------------------------------------

    slide = prs.slides.add_slide(
        prs.slide_layouts[1]
    )

    slide.shapes.title.text = (
        "Executive Summary & Recommendation"
    )

    slide.placeholders[1].text = summary

    # ------------------------------------------
    # SAVE TO MEMORY
    # ------------------------------------------

    output = io.BytesIO()

    prs.save(output)

    output.seek(0)

    return output


# ==================================================
# PPT DOWNLOAD
# ==================================================

@app.post("/download-ppt")
async def download_ppt(
    session_id: str = Form(...)
):

    try:

        if session_id not in dataframes:

            return {
                "success": False,
                "error": "Session not found. Please upload the CSV again."
            }

        df = dataframes[session_id]

        # Generate insights
        insights = generate_insights(df)

        # Generate summary using Ollama
        summary = generate_ai_summary(
            insights
        )

        # Create PPT
        ppt_file = create_ppt(
            insights,
            summary
        )

        return StreamingResponse(
            ppt_file,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation"
            ),
            headers={
                "Content-Disposition":
                "attachment; "
                "filename=CIA_Executive_Report.pptx"
            }
        )

    except Exception as e:

        return {
            "success": False,
            "error": str(e)
        }