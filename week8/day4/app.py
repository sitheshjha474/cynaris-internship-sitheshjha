
import io
import uuid
import numbers
from pathlib import Path

import pandas as pd
import plotly.express as px

from fastapi import FastAPI, UploadFile, File, Form
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pptx import Presentation


# ==================================================
# APP CONFIGURATION
# ==================================================

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"

app = FastAPI(title="CIA CSV Analyst")

# Temporary in-memory CSV storage
dataframes = {}


# ==================================================
# FRONTEND
# ==================================================

if not STATIC_DIR.is_dir():
    raise RuntimeError(f"Static folder not found: {STATIC_DIR}")

app.mount(
    "/static",
    StaticFiles(directory=str(STATIC_DIR)),
    name="static"
)


@app.get("/")
def home():
    return FileResponse(str(STATIC_DIR / "index.html"))


@app.get("/health")
def health():
    return {"status": "ok", "app": "CIA CSV Analyst"}


# ==================================================
# HELPER FUNCTIONS
# ==================================================

def format_value(value):
    if value is None:
        return "No result found."

    try:
        if pd.isna(value):
            return "No result found."
    except (TypeError, ValueError):
        pass

    if isinstance(value, numbers.Number):
        value = float(value)

        if not pd.notna(value):
            return "No result found."

        if value.is_integer():
            return f"{int(value):,}"

        return f"{value:,.2f}"

    return str(value)


def find_column(df, possible_names):
    normalized = {
        str(column).lower().replace("_", " ").strip(): column
        for column in df.columns
    }

    for name in possible_names:
        name = name.lower().replace("_", " ").strip()

        if name in normalized:
            return normalized[name]

    return None


def numeric_series(df, column):
    return pd.to_numeric(df[column], errors="coerce")


def get_numeric_column(df, question):
    q = question.lower().replace("_", " ")

    # Match an explicitly mentioned numeric column first.
    for column in df.select_dtypes(include="number").columns:
        name = str(column).lower().replace("_", " ")
        if name in q:
            return column

    # Common business metric names.
    return find_column(
        df,
        [
            "sales",
            "sales_inr",
            "sale",
            "revenue",
            "booking_value",
            "amount",
            "profit",
            "quantity",
            "price"
        ]
    )


# ==================================================
# CSV UPLOAD
# ==================================================

@app.post("/upload")
async def upload_csv(file: UploadFile = File(...)):
    try:
        filename = file.filename or ""

        if not filename.lower().endswith(".csv"):
            return {
                "success": False,
                "error": "Please upload a CSV file."
            }

        contents = await file.read()

        if not contents:
            return {
                "success": False,
                "error": "The uploaded CSV file is empty."
            }

        try:
            df = pd.read_csv(io.BytesIO(contents))
        except UnicodeDecodeError:
            df = pd.read_csv(
                io.BytesIO(contents),
                encoding="latin-1"
            )

        if df.empty and len(df.columns) == 0:
            return {
                "success": False,
                "error": "The CSV has no usable columns."
            }

        # Remove whitespace from column names.
        df.columns = [str(c).strip() for c in df.columns]

        if len(df.columns) == 0:
            return {
                "success": False,
                "error": "No columns were found in the CSV."
            }

        session_id = str(uuid.uuid4())
        dataframes[session_id] = df

        return {
            "success": True,
            "session_id": session_id,
            "filename": filename,
            "rows": int(len(df)),
            "columns": [str(c) for c in df.columns],
            "message": "CSV uploaded successfully."
        }

    except Exception as e:
        return {
            "success": False,
            "error": str(e)
        }


# ==================================================
# FREE RULE-BASED QUESTION ANSWERING
# ==================================================

def answer_question(df, question):
    q = question.lower().replace("_", " ").strip()

    sales_col = find_column(
        df,
        [
            "sales", "sales_inr", "sale", "revenue",
            "amount", "booking_value", "sales amount"
        ]
    )

    profit_col = find_column(
        df, ["profit", "profit_inr", "net profit"]
    )

    category_col = find_column(
        df, ["category", "product_category"]
    )

    city_col = find_column(
        df, ["city"]
    )

    payment_col = find_column(
        df, ["payment_method", "payment"]
    )

    numeric_cols = df.select_dtypes(
        include="number"
    ).columns.tolist()

    # Dataset size
    if any(x in q for x in [
        "how many rows", "number of rows",
        "how many records", "record count",
        "dataset size", "row count"
    ]):
        return f"The dataset contains {len(df):,} records."

    if "how many columns" in q or "number of columns" in q:
        return (
            f"The dataset contains {len(df.columns)} columns: "
            + ", ".join(map(str, df.columns))
        )

    # Show columns
    if "column names" in q or "list columns" in q:
        return "Available columns: " + ", ".join(
            map(str, df.columns)
        )

    # Missing values
    if any(x in q for x in [
        "missing values", "null values",
        "missing data", "nulls", "isnull"
    ]):
        missing = df.isna().sum()
        missing = missing[missing > 0]

        if missing.empty:
            return "There are no missing values in the dataset."

        return "Missing values by column:\n" + missing.to_string()

    # Duplicate records
    if "duplicate" in q:
        count = int(df.duplicated().sum())
        return f"The dataset contains {count:,} duplicate rows."

    # Unique values
    if "unique" in q or "distinct" in q:
        matched_col = next(
            (
                c for c in df.columns
                if str(c).lower().replace("_", " ") in q
            ),
            None
        )

        if matched_col is not None:
            count = int(df[matched_col].nunique(dropna=True))
            return (
                f"{matched_col} contains {count:,} unique values."
            )

        return (
            "Please include a column name, for example: "
            "'How many unique cities are there?'"
        )

    # Highest-selling category
    if category_col and "categor" in q and any(
        x in q for x in ["top", "highest", "best", "most", "sales"]
    ):
        if sales_col:
            temp = df[[category_col]].copy()
            temp["_metric"] = numeric_series(df, sales_col)

            result = (
                temp.groupby(category_col, dropna=False)["_metric"]
                .sum()
                .sort_values(ascending=False)
                .head(10)
            )

            if not result.empty:
                return (
                    f"Category sales ranked by {sales_col}:\n"
                    + result.to_string()
                )

        counts = df[category_col].value_counts().head(10)

        return "Most frequent categories:\n" + counts.to_string()

    # City analysis
    if city_col and "city" in q and any(
        x in q for x in ["top", "highest", "best", "most", "sales"]
    ):
        if sales_col:
            temp = df[[city_col]].copy()
            temp["_metric"] = numeric_series(df, sales_col)

            result = (
                temp.groupby(city_col, dropna=False)["_metric"]
                .sum()
                .sort_values(ascending=False)
                .head(10)
            )

            return (
                f"Top cities ranked by {sales_col}:\n"
                + result.to_string()
            )

        return (
            "Transactions by city:\n"
            + df[city_col].value_counts().head(10).to_string()
        )

    # Payment analysis
    if payment_col and "payment" in q:
        return (
            "Payment method transaction counts:\n"
            + df[payment_col].value_counts().head(10).to_string()
        )

    # Average
    if any(x in q for x in ["average", "mean"]):
        col = get_numeric_column(df, q)

        if col is None and numeric_cols:
            col = numeric_cols[0]

        if col is not None:
            value = numeric_series(df, col).mean()
            return f"Average {col}: {format_value(value)}."

        return "No numeric column is available to calculate an average."

    # Maximum / minimum
    if any(x in q for x in ["maximum", "highest value", "largest value", "maximum value"]):
        col = get_numeric_column(df, q)
        if col is not None:
            return f"Maximum {col}: {format_value(numeric_series(df, col).max())}."

    if any(x in q for x in ["minimum", "lowest value", "smallest value", "minimum value"]):
        col = get_numeric_column(df, q)
        if col is not None:
            return f"Minimum {col}: {format_value(numeric_series(df, col).min())}."

    # Total / sum
    if any(x in q for x in ["total", "sum", "overall sales", "overall revenue"]):
        col = get_numeric_column(df, q)

        if col is None and numeric_cols:
            col = numeric_cols[0]

        if col is not None:
            value = numeric_series(df, col).sum()
            return f"Total {col}: {format_value(value)}."

        return "No numeric column is available for a total calculation."

    # Descriptive statistics
    if any(x in q for x in [
        "describe", "descriptive statistics",
        "statistical summary", "summary statistics"
    ]):
        if numeric_cols:
            return df[numeric_cols].describe().round(2).to_string()

        return "There are no numeric columns for statistical analysis."

    # Top numeric values
    if any(x in q for x in ["top 10", "top ten", "highest values"]):
        col = get_numeric_column(df, q)

        if col is not None:
            return (
                f"Top values in {col}:\n"
                + numeric_series(df, col).nlargest(10).to_string(index=False)
            )

    return (
        "I couldn't match that question to a supported analysis. "
        "Try: 'How many rows are there?', 'What is the total sales?', "
        "'What is the average sales?', 'Show top cities by sales', "
        "'Which category has the highest sales?', "
        "'Show payment method counts', 'Are there missing values?', "
        "or 'How many duplicate rows are there?'. "
        "Available columns: " + ", ".join(map(str, df.columns))
    )


# ==================================================
# CREATE PLOTLY CHART JSON
# ==================================================

def create_chart(df, question):
    q = question.lower()

    category_col = find_column(
        df, ["category", "product_category"]
    )
    city_col = find_column(
        df, ["city"]
    )
    sales_col = find_column(
        df,
        [
            "sales", "sales_inr", "sale", "revenue",
            "amount", "booking_value", "sales amount"
        ]
    )
    payment_col = find_column(
        df, ["payment_method", "payment"]
    )

    # Category sales
    if "categor" in q and any(
        x in q for x in ["sales", "revenue", "amount", "profit"]
    ):
        if category_col and sales_col:
            chart_df = pd.DataFrame({
                "category": df[category_col],
                "value": numeric_series(df, sales_col)
            })

            chart_df = (
                chart_df.groupby("category", dropna=False)["value"]
                .sum()
                .sort_values(ascending=False)
                .head(10)
                .reset_index()
            )

            fig = px.bar(
                chart_df,
                x="category",
                y="value",
                title="Sales by Category",
                labels={
                    "category": "Category",
                    "value": sales_col
                }
            )

            return fig.to_json()

    # City sales
    if "city" in q and any(
        x in q for x in ["sales", "revenue", "amount", "profit"]
    ):
        if city_col and sales_col:
            chart_df = pd.DataFrame({
                "city": df[city_col],
                "value": numeric_series(df, sales_col)
            })

            chart_df = (
                chart_df.groupby("city", dropna=False)["value"]
                .sum()
                .sort_values(ascending=False)
                .head(10)
                .reset_index()
            )

            fig = px.bar(
                chart_df,
                x="city",
                y="value",
                title="Top Cities by Sales",
                labels={"city": "City", "value": sales_col}
            )

            return fig.to_json()

    # Payment methods
    if "payment" in q and payment_col:
        chart_df = (
            df[payment_col]
            .value_counts()
            .rename_axis("payment_method")
            .reset_index(name="count")
        )

        fig = px.bar(
            chart_df,
            x="payment_method",
            y="count",
            title="Payment Method Usage"
        )

        return fig.to_json()

    # Generic numeric distribution
    if any(x in q for x in ["distribution", "histogram"]):
        col = get_numeric_column(df, q)

        if col is not None:
            fig = px.histogram(
                df,
                x=col,
                title=f"Distribution of {col}"
            )

            return fig.to_json()

    return None


# ==================================================
# QUESTION ENDPOINT
# ==================================================

@app.post("/cia/sql-analyst")
async def cia_query(
    session_id: str = Form(...),
    question: str = Form(...)
):
    try:
        if session_id not in dataframes:
            return {
                "success": False,
                "error": "Session not found. Please upload the CSV again."
            }

        df = dataframes[session_id]

        answer = answer_question(df, question)
        chart = create_chart(df, question)

        return {
            "success": True,
            "question": question,
            "answer": answer,
            "chart": chart
        }

    except Exception as e:
        return {
            "success": False,
            "error": str(e)
        }


# ==================================================
# GENERATE FIVE EDA INSIGHTS
# ==================================================

def generate_insights(df):
    insights = []

    sales_col = find_column(
        df,
        [
            "sales", "sales_inr", "sale", "revenue",
            "amount", "booking_value", "sales amount"
        ]
    )

    profit_col = find_column(
        df, ["profit", "profit_inr", "net profit"]
    )

    category_col = find_column(
        df, ["category", "product_category"]
    )

    city_col = find_column(
        df, ["city"]
    )

    payment_col = find_column(
        df, ["payment_method", "payment"]
    )

    # 1. Dataset size
    insights.append(
        f"The dataset contains {len(df):,} records "
        f"and {len(df.columns)} columns."
    )

    # 2. Total sales
    if sales_col:
        total_sales = numeric_series(df, sales_col).sum()

        insights.append(
            f"Total {sales_col} is {format_value(total_sales)}."
        )
    else:
        insights.append(
            "A recognized sales or revenue column was not found."
        )

    # 3. Highest-selling category
    if category_col and sales_col:
        temp = pd.DataFrame({
            "category": df[category_col],
            "sales": numeric_series(df, sales_col)
        })

        totals = temp.groupby(
            "category", dropna=False
        )["sales"].sum().sort_values(ascending=False)

        if not totals.empty:
            insights.append(
                f"{totals.index[0]} is the highest-selling category "
                f"with {format_value(totals.iloc[0])} in {sales_col}."
            )
        else:
            insights.append("No category sales results were available.")
    else:
        insights.append(
            "Category sales analysis was unavailable because "
            "a category or sales column was not found."
        )

    # 4. Highest-performing city
    if city_col and sales_col:
        temp = pd.DataFrame({
            "city": df[city_col],
            "sales": numeric_series(df, sales_col)
        })

        totals = temp.groupby(
            "city", dropna=False
        )["sales"].sum().sort_values(ascending=False)

        if not totals.empty:
            insights.append(
                f"{totals.index[0]} is the highest-performing city "
                f"with {format_value(totals.iloc[0])} in {sales_col}."
            )
        else:
            insights.append("No city sales results were available.")
    else:
        insights.append(
            "City sales analysis was unavailable because "
            "a city or sales column was not found."
        )

    # 5. Average profit or most-used payment method
    if profit_col:
        average_profit = numeric_series(df, profit_col).mean()

        insights.append(
            f"Average {profit_col} is {format_value(average_profit)}."
        )

    elif payment_col:
        counts = df[payment_col].value_counts()

        if not counts.empty:
            insights.append(
                f"{counts.index[0]} is the most frequently used "
                f"payment method with {int(counts.iloc[0]):,} records."
            )
        else:
            insights.append("No payment method results were available.")

    else:
        insights.append(
            "Profit and payment method columns were not found."
        )

    return insights[:5]


# ==================================================
# FREE EXECUTIVE SUMMARY
# ==================================================

def generate_ai_summary(insights):
    # No external AI model is used. This is a factual,
    # template-based summary derived from the five insights.

    usable = [
        str(item) for item in insights
        if item
    ]

    if not usable:
        return (
            "The dataset was processed, but no meaningful insights "
            "were available. Recommendation: review the data structure "
            "and ensure the relevant business columns are present."
        )

    summary = (
        "The analysis provides an overview of the dataset and its "
        "available business performance measures. "
        + " ".join(usable)
        + " Recommendation: validate data quality, investigate the "
        "strongest-performing categories and locations where available, "
        "and use these findings to guide further business analysis."
    )

    words = summary.split()

    if len(words) > 200:
        summary = " ".join(words[:200])

    return summary


# ==================================================
# AI SUMMARY ENDPOINT (FREE TEMPLATE SUMMARY)
# ==================================================

@app.post("/ai-summary")
async def ai_summary(
    session_id: str = Form(...)
):
    try:
        if session_id not in dataframes:
            return {
                "success": False,
                "error": "Session not found. Please upload the CSV again."
            }

        df = dataframes[session_id]
        insights = generate_insights(df)
        summary = generate_ai_summary(insights)

        return {
            "success": True,
            "insights": insights,
            "summary": summary
        }

    except Exception as e:
        return {
            "success": False,
            "error": str(e)
        }


# ==================================================
# CREATE POWERPOINT REPORT
# ==================================================

def create_ppt(insights, summary):
    prs = Presentation()

    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "CIA Data Analysis Report"

    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = (
            "CSV Analysis and Executive Summary"
        )

    # Slides 2-6: Five insights
    for index, insight in enumerate(insights, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Key Insight {index}"
        slide.placeholders[1].text = str(insight)

    # Slide 7: Summary and recommendation
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Executive Summary & Recommendation"
    slide.placeholders[1].text = str(summary)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)

    return output


# ==================================================
# POWERPOINT DOWNLOAD
# ==================================================

@app.post("/download-ppt")
async def download_ppt(
    session_id: str = Form(...)
):
    try:
        if session_id not in dataframes:
            return {
                "success": False,
                "error": "Session not found. Please upload the CSV again."
            }

        df = dataframes[session_id]
        insights = generate_insights(df)
        summary = generate_ai_summary(insights)

        ppt_file = create_ppt(insights, summary)

        return StreamingResponse(
            ppt_file,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "presentationml.presentation"
            ),
            headers={
                "Content-Disposition":
                    'attachment; filename="CIA_Executive_Report.pptx"'
            }
        )

    except Exception as e:
        return {
            "success": False,
            "error": str(e)
        }
